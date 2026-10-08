"""Build the unit database (Excel + web JSON) from the reconciled brochure extractions.

Inputs  : pipeline/raw/<project>-{dist,floor,site}.json  (reconciled agent extractions)
Outputs : data/BTO_Units_June2026.xlsx   (tabs: Plus, Prime, Standard, Facilities, Method)
          data/units/<project>.json, data/units/index.json   (web app)

Every unit row is generated from the brochure's Unit Distribution grid; all other
attributes are joined on (block, unit stack[, storey]). Any unit that cannot be joined
aborts the build -- nothing is silently defaulted.
"""
import json
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "pipeline" / "raw"
OUT_UNITS = ROOT / "data" / "units"

# Thresholds (metres). Change here and rebuild.
LIFT_NEAR_M = 15
CHUTE_NEAR_M = 10
MRT_NEAR_M = 400
FACILITY_NEAR_M = 50
CLEARANCE_M = 30

PROJECTS = [
    # key, name, town, type, short code
    ("kebun-baru_breeze-5", "Kebun Baru Breeze", "Ang Mo Kio", "Plus", "KBB"),
    ("kebun-baru_ridge-1", "Kebun Baru Ridge", "Ang Mo Kio", "Plus", "KBR"),
    ("berlayar-rise-7", "Berlayar Rise", "Bukit Merah", "Prime", "BR"),
    ("lakeview-cascadia-2", "Lakeview Cascadia", "Bishan", "Prime", "LC"),
    ("sembawang-brook-6", "Sembawang Brook", "Sembawang", "Standard", "SB"),
    ("sembawang-portico-4", "Sembawang Portico", "Sembawang", "Standard", "SP"),
    ("woodgrove-acres-3", "Woodgrove Acres", "Woodlands", "Standard", "WA"),
]

COLUMNS = [
    ("id", "Unit ID"), ("project", "Project"), ("project_type", "Project type"), ("town", "Town"),
    ("block", "Block"), ("storey", "Storey"), ("unit", "Unit no."), ("address", "Unit (#storey-unit)"),
    ("flat_type", "Unit type"), ("unit_design", "Unit design (2RF)"),
    ("facing", "Sun direction (facing)"), ("facing_bearing", "Facing bearing (deg)"),
    ("position", "Corridor / Corner"),
    ("lift_m", "Lift distance (m)"), ("lift_near", "Near lift (<=15 m)"),
    ("chute_m", "Rubbish chute distance (m)"), ("chute_near", "Near rubbish chute (<=10 m)"),
    ("roof_access", "Access to roof"),
    ("neighbour_m", "Distance to nearest other block (m)"), ("gt30m", ">30 m from neighbouring blocks"),
    ("mrt_name", "Nearest MRT"), ("mrt_m", "MRT distance (m)"), ("mrt_near", "Near MRT (<=400 m)"), ("mrt_estimated", "MRT distance estimated"),
    ("fac_has_count", "Facilities in block (count)"), ("facilities_has", "Facilities in block"),
    ("fac_near_count", "Facilities nearby (count)"), ("facilities_near", "Facilities nearby (<=50 m)"),
    ("source", "Source pages"),
]


def load(key, kind):
    path = RAW / f"{key}-{kind}.json"
    if not path.exists():
        sys.exit(f"missing {path}")
    return json.loads(path.read_text())


def norm_unit(u):
    return str(u).strip().lstrip("#").lstrip("0") or "0"


def norm_block(b):
    return str(b).upper().replace("BLOCK", "").replace("BLK", "").strip()


def bearing_to_8pt(b):
    return ["N", "NE", "E", "SE", "S", "SW", "W", "NW"][int(((b % 360) + 22.5) // 45) % 8]


def build_project(key, name, town, ptype, code, errors):
    dist, floor, site = load(key, "dist"), load(key, "floor"), load(key, "site")
    mpp = site["scale"]["m_per_px"]

    # site geometry per stack
    pos = {}
    for s in site["stacks"]:
        k = (norm_block(s["block"]), norm_unit(s["unit"]))
        if k in pos:
            errors.append(f"{key}: duplicate site stack {k}")
        pos[k] = s

    def dist_m(a, b):
        return math.hypot(a["x"] - b["x"], a["y"] - b["y"]) * mpp

    # floor-plan attributes per (block, unit, storey)
    fp = {}
    for s in floor["stacks"]:
        for st in s["storeys"]:
            k = (norm_block(s["block"]), norm_unit(s["unit"]), st)
            if k in fp:
                errors.append(f"{key}: floor-plan attributes defined twice for {k}")
            fp[k] = s

    on_site_mrt = [m for m in site["mrt"] if m.get("on_site_plan") and m.get("x") is not None]
    off_site_mrt = [m for m in site["mrt"] if not m.get("on_site_plan") and m.get("approx_distance_m_to_site_centre")]

    units, seen = [], set()
    for e in dist["entries"]:
        block, unit = norm_block(e["block"]), norm_unit(e["unit"])
        sp = pos.get((block, unit))
        if sp is None:
            errors.append(f"{key}: no site-plan position for block {block} unit {unit}")
            continue
        others = [s for (b, _), s in pos.items() if b != block]
        neighbour = min((dist_m(sp, o) for o in others), default=None)
        has = [f["name"] for f in site["facilities"] if f.get("located_in_block") and norm_block(f["located_in_block"]) == block]
        near = sorted({f["name"] for f in site["facilities"]
                       if not f.get("located_in_block") and dist_m(sp, f) <= FACILITY_NEAR_M})
        if on_site_mrt:
            m = min(on_site_mrt, key=lambda m: dist_m(sp, m))
            mrt_name, mrt_m, mrt_est = m["name"], dist_m(sp, m), False
        elif off_site_mrt:
            m = min(off_site_mrt, key=lambda m: m["approx_distance_m_to_site_centre"])
            mrt_name, mrt_m, mrt_est = m["name"], m["approx_distance_m_to_site_centre"], True
        else:
            mrt_name, mrt_m, mrt_est = None, None, True
        for st in sorted(e["storeys"]):
            uid = f"{code}-{block}-{st:02d}-{unit}"
            if uid in seen:
                errors.append(f"{key}: duplicate unit {uid}")
                continue
            seen.add(uid)
            f = fp.get((block, unit, st))
            if f is None:
                errors.append(f"{key}: no floor-plan attributes for block {block} #{st:02d}-{unit}")
                continue
            design = None
            if e["flat_type"].startswith("2-Room Flexi"):
                design = "Type 1" if "Type 1" in e["flat_type"] else "Type 2"
            units.append({
                "id": uid, "project": name, "project_key": key, "project_type": ptype, "town": town,
                "block": block, "storey": st, "unit": unit, "address": f"Blk {block} #{st:02d}-{unit}",
                "flat_type": e["flat_type"], "unit_design": design,
                "facing": sp["facing_8pt"], "facing_bearing": round(sp["facing_bearing_deg"]),
                "position": f["position"].capitalize(),
                "lift_m": f["lift_distance_m"], "lift_near": f["lift_distance_m"] is not None and f["lift_distance_m"] <= LIFT_NEAR_M,
                "chute_m": f["chute_distance_m"], "chute_near": f["chute_distance_m"] is not None and f["chute_distance_m"] <= CHUTE_NEAR_M,
                "roof_access": st in f["roof_access_floors"],
                "neighbour_m": round(neighbour, 1) if neighbour is not None else None,
                "gt30m": neighbour is None or neighbour > CLEARANCE_M,
                "mrt_name": mrt_name, "mrt_m": round(mrt_m) if mrt_m is not None else None,
                "mrt_near": mrt_m is not None and mrt_m <= MRT_NEAR_M, "mrt_estimated": mrt_est,
                "fac_has_count": len(has), "facilities_has": has,
                "fac_near_count": len(near), "facilities_near": near,
                "source": f"dist p{','.join(map(str, dist['source_pages']))}; plan p{','.join(map(str, f['floorplan_pages']))}; site p{site['site_plan_page']}",
            })
    units.sort(key=lambda u: (u["block"], u["storey"], u["unit"]))
    return units, site


def write_excel(all_units, sites):
    wb = Workbook()
    wb.remove(wb.active)
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for ptype in ("Plus", "Prime", "Standard"):
        ws = wb.create_sheet(ptype)
        ws.append([label for _, label in COLUMNS])
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
        for u in (u for u in all_units if u["project_type"] == ptype):
            row = []
            for k, _ in COLUMNS:
                v = u[k]
                if isinstance(v, list):
                    v = "; ".join(v)
                elif isinstance(v, bool):
                    v = "Yes" if v else "No"
                row.append(v)
            ws.append(row)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, (k, label) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(label) + 2))
    fs = wb.create_sheet("Facilities")
    fs.append(["Project", "Legend no.", "Facility", "In residential block", "Host", "Site plan x (px)", "Site plan y (px)"])
    for key, site in sites.items():
        for f in site["facilities"]:
            fs.append([key, f["legend_no"], f["name"], f.get("located_in_block"), f.get("host"), f["x"], f["y"]])
    ms = wb.create_sheet("Method")
    for line in [
        "Source: HDB June 2026 BTO sales brochures (PDF). One row per sale unit from each block's Unit Distribution grid (rental / non-sale cells excluded).",
        "Sun direction: compass direction the living/dining-room windows face (8-point), from the site plan north arrow and the block floor plan.",
        "Corridor/Corner: Corner = stack at the outer end of a row of units (windows on two external faces).",
        f"Lift / rubbish chute distance: straight-line metres from the unit's main door to the lift lobby / refuse chute, measured on the floor plan scale bar. Near lift <= {LIFT_NEAR_M} m, near chute <= {CHUTE_NEAR_M} m.",
        "Access to roof: unit adjoins/overlooks a roof annotated on the floor plan (e.g. 'ROOF AT 2ND STOREY ONLY') on that storey.",
        f"> {CLEARANCE_M} m from neighbouring blocks: distance from the unit stack's centre to the nearest stack centre of any other block, measured on the site plan scale bar.",
        f"Facilities: 'in block' = facility located in/at/on the unit's own residential block; 'nearby' = facility not in any residential block whose site-plan marker is within {FACILITY_NEAR_M} m of the unit stack's centre.",
        f"MRT: straight-line distance from the stack centre to the station on the site plan (or from the site centre on the location plan when the station is off the site plan; flagged as estimated). Near <= {MRT_NEAR_M} m.",
    ]:
        ms.append([line])
    ms.column_dimensions["A"].width = 160
    path = ROOT / "data" / "BTO_Units_June2026.xlsx"
    wb.save(path)
    return path


def main():
    errors, all_units, sites, index = [], [], {}, []
    OUT_UNITS.mkdir(parents=True, exist_ok=True)
    for key, name, town, ptype, code in PROJECTS:
        units, site = build_project(key, name, town, ptype, code, errors)
        sites[key] = site
        all_units += units
        (OUT_UNITS / f"{key}.json").write_text(json.dumps(units, separators=(",", ":")))
        index.append({"key": key, "name": name, "town": town, "project_type": ptype, "units": len(units),
                      "blocks": sorted({u["block"] for u in units}), "max_storey": max((u["storey"] for u in units), default=0),
                      "flat_types": sorted({u["flat_type"] for u in units})})
    if errors:
        print("BUILD FAILED:\n  " + "\n  ".join(errors))
        sys.exit(1)
    flat_types = sorted({u["flat_type"] for u in all_units})
    (OUT_UNITS / "index.json").write_text(json.dumps({"projects": index, "flat_types": flat_types}, indent=1))
    path = write_excel(all_units, sites)
    for p in index:
        print(f"{p['project_type']:9} {p['name']:20} {p['units']:5} units  blocks {', '.join(p['blocks'])}")
    print(f"total {len(all_units)} units -> {path}")


if __name__ == "__main__":
    main()
