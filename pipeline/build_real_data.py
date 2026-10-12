"""Build real unit data from the per-project analysis files pipeline/raw/real/<project_key>.json.

Default: validate every file and write the review workbook review/mybto-units-review.xlsx and
review/mybto-units-review.csv (review/ is git-ignored, the workbook is shared with the user directly).
With --publish: also write public/data/units/<key>.json and index.json in the schema the site uses
(same fields as make_unit in make_sample_data.py). A project whose real file is missing keeps its current
sample file and stays in index.json.

Dev-only dependency: openpyxl (pip install openpyxl). Not part of package.json.

  python3 pipeline/build_real_data.py                      # validate + review workbook
  python3 pipeline/build_real_data.py --publish            # also write public/data/units (after user approval)
  python3 pipeline/build_real_data.py --in DIR --out DIR   # alternative input folder / output root (tests)
"""
import argparse
import csv
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
import make_sample_data as msd  # noqa: E402  (PROJECTS, HOUSING, thresholds; importing only defines things)

FACINGS = {"NE", "SE", "SW", "NW"}
FLAT_TYPES = ["2-Room Flexi (Type 1)", "2-Room Flexi (Type 2)", "3-Room", "4-Room", "5-Room", "3Gen"]
PROJ = {p[0]: p for p in msd.PROJECTS}  # key -> (key, name, town, type, code, blocks, types, mrt)


def storey_ranges(storeys):
    """[2,3,4,6] -> '02-04,06'."""
    out, s = [], sorted(set(storeys))
    i = 0
    while i < len(s):
        j = i
        while j + 1 < len(s) and s[j + 1] == s[j] + 1:
            j += 1
        out.append(f"{s[i]:02d}" if i == j else f"{s[i]:02d}-{s[j]:02d}")
        i = j + 1
    return ",".join(out)


def validate(key, d):
    """Return (errors, warnings) lists of strings."""
    err, warn = [], []
    stacks = {}
    for s in d.get("stacks", []):
        k = (s.get("block"), s.get("unit"))
        if k in stacks:
            err.append(f"duplicate stack {k}")
        stacks[k] = s
        if s.get("facing") not in FACINGS:
            err.append(f"stack {k}: facing {s.get('facing')!r} not in NE/SE/SW/NW")
        if s.get("gt30m") is None:
            warn.append(f"stack {k}: gt30m is null")
        if s.get("position") not in ("Corner", "Corridor"):
            warn.append(f"stack {k}: position {s.get('position')!r}")
        if s.get("chute_m") is None:
            warn.append(f"stack {k}: chute_m is null")
        if s.get("wind_open") is None:
            warn.append(f"stack {k}: wind_open is null")
        if s.get("confidence") not in ("high", None):
            warn.append(f"stack {k}: confidence {s.get('confidence')}")
    seen = set()
    for u in d.get("units", []):
        uk = (u.get("block"), u.get("storey"), u.get("unit"))
        if uk in seen:
            err.append(f"duplicate unit {uk}")
        seen.add(uk)
        if (u.get("block"), u.get("unit")) not in stacks:
            err.append(f"unit {uk}: stack missing")
        if u.get("flat_type") not in FLAT_TYPES:
            err.append(f"unit {uk}: flat_type {u.get('flat_type')!r} not allowed")
        if not isinstance(u.get("storey"), int):
            err.append(f"unit {uk}: storey is not an int")
    unused = set(stacks) - {(u.get("block"), u.get("unit")) for u in d.get("units", [])}
    for k in sorted(unused, key=str):
        warn.append(f"stack {k} has no units")
    if not d.get("units"):
        err.append("no units")
    if not (d.get("totals_check") or {}).get("match"):
        err.append("totals_check.match is not true")
    return err, warn


def load_all(in_dir):
    """Return {key: data} for the known projects; missing or broken files are skipped with a warning."""
    data = {}
    for key in PROJ:
        path = in_dir / f"{key}.json"
        if not path.exists():
            print(f"WARNING {key}: {path.name} missing, skipped (sample data stays)")
            continue
        try:
            d = json.loads(path.read_text())
        except ValueError as e:
            print(f"WARNING {key}: invalid JSON ({e}), skipped")
            continue
        err, warn = validate(key, d)
        for w in warn:
            print(f"WARNING {key}: {w}")
        for e in err:
            print(f"ERROR   {key}: {e}")
        d["_errors"] = err
        data[key] = d
    return data


def unit_rows(key, d):
    """One dict per unit, in the website schema (plus mscp_estimated) built from stack attributes."""
    _, name, town, ptype, code, _b, _t, mrt = PROJ[key]
    stacks = {(s["block"], s["unit"]): s for s in d["stacks"]}
    rows = []
    for u in d["units"]:
        s = stacks.get((u["block"], u["unit"]))
        if s is None:
            continue
        storey, ft = u["storey"], u["flat_type"]
        design = None
        if ft.startswith("2-Room Flexi"):
            design = "Type 1" if "Type 1" in ft else "Type 2"
        gt = s.get("gt30m")
        chute, mrt_m, mscp = s.get("chute_m"), s.get("mrt_m"), s.get("mscp_m")
        has, near = list(s.get("facilities_has") or []), list(s.get("facilities_near") or [])
        rows.append({
            "id": f"{code}-{u['block']}-{storey:02d}-{u['unit']}", "project": name, "project_key": key,
            "project_type": ptype, "town": town, "block": u["block"], "storey": storey, "unit": u["unit"],
            "address": f"Blk {u['block']} #{storey:02d}-{u['unit']}", "flat_type": ft, "unit_design": design,
            "facing": s["facing"], "position": s.get("position"),
            "lift_m": None, "lift_near": False,
            "chute_m": chute, "chute_near": chute is not None and chute <= msd.CHUTE_NEAR_M,
            "roof_access": storey in (s.get("roof_floors") or []),
            "neighbour_m": None, "gt30m": gt,
            "opposite_m": None, "opposite_gt30m": gt,
            "mrt_name": mrt, "mrt_m": mrt_m, "mrt_near": mrt_m is not None and mrt_m <= msd.MRT_NEAR_M,
            "wind_open": s.get("wind_open"),
            "mscp_m": mscp, "mscp_near": mscp is not None and mscp <= msd.MSCP_NEAR_M,
            "mscp_estimated": mscp is not None,
            "facilities_has": has, "fac_has_count": len(has),
            "facilities_near": near, "fac_near_count": len(near),
            "sample_fields": [],
        })
    rows.sort(key=lambda r: (r["block"], r["storey"], r["unit"]))
    return rows


# ---------------------------------------------------------------- review workbook

def summary_row(key, d, rows):
    stacks = d["stacks"]
    tc = d.get("totals_check") or {}
    brochure, counted = tc.get("brochure") or {}, tc.get("counted") or {}
    types = [t for t in FLAT_TYPES if t in brochure or t in counted or any(r["flat_type"] == t for r in rows)]
    per_type = "; ".join(f"{t}: {sum(1 for r in rows if r['flat_type'] == t)} / {brochure.get(t, '?')}" for t in types)
    n = lambda f: sum(1 for r in rows if f(r))  # noqa: E731
    return {
        "Project": d.get("name") or PROJ[key][1], "Key": key, "Stacks": len(stacks), "Units": len(rows),
        "Flat types (counted / brochure)": per_type, "Totals match": "PASS" if tc.get("match") else "FAIL",
        "gt30m pass": sum(1 for s in stacks if s.get("gt30m") is True),
        "gt30m fail": sum(1 for s in stacks if s.get("gt30m") is False),
        "gt30m unknown": sum(1 for s in stacks if s.get("gt30m") is None),
        "Corner stacks": sum(1 for s in stacks if s.get("position") == "Corner"),
        "Chute <=10 m stacks": sum(1 for s in stacks if s.get("chute_m") is not None and s["chute_m"] <= msd.CHUTE_NEAR_M),
        "wind_open stacks": sum(1 for s in stacks if s.get("wind_open")),
        "MSCP within 100 m stacks": sum(1 for s in stacks if s.get("mscp_m") is not None and s["mscp_m"] <= msd.MSCP_NEAR_M),
        "MRT within 100 m stacks": sum(1 for s in stacks if s.get("mrt_m") is not None and s["mrt_m"] <= msd.MRT_NEAR_M),
        "Low-confidence stacks": ", ".join(f"{s['block']}-{s['unit']}" for s in stacks if s.get("confidence") not in ("high", None)),
        "Errors": "; ".join(d.get("_errors") or []),
        "Notes": " | ".join(d.get("notes") or []),
    }


STACK_COLS = ["Block", "Unit", "Storeys", "Flat types", "Facing", "gt30m", "gt30m evidence crop", "Position", "Chute m",
              "Roof floors", "wind_open", "MSCP m", "MRT m", "Facilities (has)", "Facilities (near)", "Confidence", "Notes"]


def stack_row(s, units):
    mine = [u for u in units if u["block"] == s["block"] and u["unit"] == s["unit"]]
    ev = s.get("gt30m_evidence") or {}
    return [s["block"], s["unit"], storey_ranges(u["storey"] for u in mine),
            ", ".join(t for t in FLAT_TYPES if any(u["flat_type"] == t for u in mine)),
            s.get("facing"), s.get("gt30m"), ev.get("crop"), s.get("position"), s.get("chute_m"),
            ",".join(str(x) for x in s.get("roof_floors") or []), s.get("wind_open"), s.get("mscp_m"), s.get("mrt_m"),
            ", ".join(s.get("facilities_has") or []), ", ".join(s.get("facilities_near") or []),
            s.get("confidence"), s.get("notes")]


ALL_COLS = ["id", "project", "block", "storey", "unit", "flat_type", "facing", "position", "gt30m", "chute_m", "chute_near",
            "roof_access", "wind_open", "mscp_m", "mscp_near", "mrt_name", "mrt_m", "mrt_near", "facilities_has",
            "facilities_near"]


def all_row(r):
    out = []
    for c in ALL_COLS:
        v = r[c]
        out.append(", ".join(v) if isinstance(v, list) else v)
    return out


def write_review(data, review_dir):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("ERROR openpyxl missing: pip install openpyxl (dev-only tool)")
        return False
    review_dir.mkdir(parents=True, exist_ok=True)
    green, red = PatternFill("solid", fgColor="D9F2D9"), PatternFill("solid", fgColor="F8D7D7")
    head = PatternFill("solid", fgColor="E8EDF5")

    def finish(ws, ncols, bool_cols=(), widths=None, wrap_cols=()):
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), head
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"
        for i in range(1, ncols + 1):
            letter = get_column_letter(i)
            vals = [len(str(c.value)) for c in ws[letter][:200] if c.value is not None]
            w = (widths or {}).get(i) or min(max(vals + [8]) + 2, 60)
            ws.column_dimensions[letter].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.column in bool_cols:
                    if c.value is True or c.value == "PASS":
                        c.fill = green
                    elif c.value is False or c.value == "FAIL":
                        c.fill = red
                if c.column in wrap_cols:
                    c.alignment = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    srows, allrows = [], []
    for key, d in data.items():
        rows = unit_rows(key, d)
        srows.append(summary_row(key, d, rows))
        allrows += rows
    cols = list(srows[0].keys()) if srows else ["Project"]
    ws.append(cols)
    for r in srows:
        ws.append([r[c] for c in cols])
    finish(ws, len(cols), bool_cols=(cols.index("Totals match") + 1,),
           widths={cols.index("Flat types (counted / brochure)") + 1: 50, cols.index("Notes") + 1: 80, cols.index("Errors") + 1: 40},
           wrap_cols=(cols.index("Flat types (counted / brochure)") + 1, cols.index("Notes") + 1, cols.index("Errors") + 1))
    for key, d in data.items():
        sh = wb.create_sheet(key[:31])
        sh.append(STACK_COLS)
        for s in sorted(d["stacks"], key=lambda s: (s["block"], s["unit"])):
            sh.append(stack_row(s, d["units"]))
        finish(sh, len(STACK_COLS), bool_cols=(STACK_COLS.index("gt30m") + 1,), widths={STACK_COLS.index("Notes") + 1: 60},
               wrap_cols=(STACK_COLS.index("Notes") + 1,))
    wa = wb.create_sheet("All units")
    wa.append(ALL_COLS)
    for r in allrows:
        wa.append(all_row(r))
    finish(wa, len(ALL_COLS), bool_cols=(ALL_COLS.index("gt30m") + 1,))
    wb.save(review_dir / "mybto-units-review.xlsx")
    with open(review_dir / "mybto-units-review.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(ALL_COLS)
        for r in allrows:
            w.writerow(all_row(r))
    print(f"wrote {review_dir / 'mybto-units-review.xlsx'} and .csv ({len(allrows)} units)")
    return True


# ---------------------------------------------------------------- publish

def publish(data, out_dir):
    units_dir = out_dir / "public" / "data" / "units"
    units_dir.mkdir(parents=True, exist_ok=True)
    idx_path = units_dir / "index.json"
    old = json.loads(idx_path.read_text()) if idx_path.exists() else {"projects": []}
    entries = {p["key"]: p for p in old.get("projects", [])}
    for key, d in data.items():
        if d["_errors"]:
            print(f"SKIP publish {key}: validation errors")
            continue
        rows = unit_rows(key, d)
        (units_dir / f"{key}.json").write_text(json.dumps(rows, separators=(",", ":")))
        _, name, town, ptype, *_ = PROJ[key]
        entries[key] = {"key": key, "name": name, "town": town, "project_type": ptype, "units": len(rows),
                        "blocks": sorted({r["block"] for r in rows}), "max_storey": max(r["storey"] for r in rows),
                        "flat_types": sorted({r["flat_type"] for r in rows}),
                        "housing_types": [h for h in msd.HOUSING if any(msd.housing(r["flat_type"]) == h for r in rows)],
                        "data_status": "real data"}
        print(f"published {key}: {len(rows)} units")
    projects = [entries[k] for k in PROJ if k in entries]  # keep the site's project order
    all_types = sorted({t for p in projects for t in p["flat_types"]})
    idx = {"sample": any(p["data_status"] != "real data" for p in projects), "projects": projects, "flat_types": all_types,
           "housing_types": [h for h in msd.HOUSING if any(msd.housing(t) == h for t in all_types)]}
    idx_path.write_text(json.dumps(idx, indent=1))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="in_dir", default=str(ROOT / "pipeline" / "raw" / "real"), help="folder with <key>.json analysis files")
    ap.add_argument("--out", dest="out_dir", default=str(ROOT), help="output root (review/ and public/data/units/ go below it)")
    ap.add_argument("--publish", action="store_true", help="also write public/data/units (only after the user approves the review)")
    a = ap.parse_args()
    out = Path(a.out_dir)
    data = load_all(Path(a.in_dir))
    if not data:
        print("no project files found, nothing to do")
        return 1
    ok = write_review(data, out / "review")
    if a.publish:
        publish(data, out)
    bad = [k for k, d in data.items() if d["_errors"]]
    if bad:
        print("validation errors in: " + ", ".join(bad))
    return 0 if ok and not bad else 1


if __name__ == "__main__":
    sys.exit(main())
